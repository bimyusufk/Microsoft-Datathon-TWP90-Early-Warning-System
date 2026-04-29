"""
Script eksplorasi struktur sheet Perbankan dari XLSX.
"""
from openpyxl import load_workbook

XLSX_PATH = r"c:\Users\bimyu\Documents\Projects\DatathonMETC\data prep datathon dicoding.xlsx"

wb = load_workbook(XLSX_PATH, data_only=True, read_only=True)
print("Sheet names:", wb.sheetnames)
print()

ws = wb["Perbankan"]
rows = list(ws.iter_rows(values_only=True))
print(f"Total rows: {len(rows)}")
print(f"Total cols: {max(len(r) for r in rows[:10]) if rows else 0}")
print()

print("=== ROW 1 (header 1) ===")
for i, v in enumerate(rows[0]):
    if v is not None:
        print(f"  col[{i}] = {repr(v)}")

print()
print("=== ROW 2 (header 2) ===")
for i, v in enumerate(rows[1]):
    if v is not None:
        print(f"  col[{i}] = {repr(v)}")

print()
print("=== ROW 3 (header 3, if any) ===")
for i, v in enumerate(rows[2]):
    if v is not None:
        print(f"  col[{i}] = {repr(v)}")

print()
print("=== DATA ROWS 4-8 (first 5 data rows) ===")
for row in rows[3:8]:
    print(row)

print()
print("=== LAST 3 ROWS ===")
for row in rows[-3:]:
    print(row)

wb.close()
