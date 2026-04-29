"""
Eksplorasi mapping ID_Pemda dari sheet Perbankan vs dim_pemda di PostgreSQL.
"""
import psycopg
from openpyxl import load_workbook

DB_URL    = "postgresql://postgres:root@localhost:5432/datathon"
XLSX_PATH = r"c:\Users\bimyu\Documents\Projects\DatathonMETC\data prep datathon dicoding.xlsx"

# Baca ID_Pemda dari sheet Perbankan
wb = load_workbook(XLSX_PATH, data_only=True, read_only=True)
ws = wb["Perbankan"]
rows = [r for r in ws.iter_rows(values_only=True)][2:]  # skip 2 header rows
wb.close()

sheet_pemda = {}
for r in rows:
    if r and r[1] not in (None, ""):
        pid_raw = r[0]
        prov    = r[1]
        if pid_raw is not None:
            try:
                pid = int(float(str(pid_raw)))
            except:
                pid = str(pid_raw).strip()
            sheet_pemda[pid] = prov

print(f"ID_Pemda unik di sheet Perbankan: {len(sheet_pemda)}")
for k, v in sorted(sheet_pemda.items(), key=lambda x: str(x[0])):
    print(f"  {str(k):>6}  →  {v}")

# Cek vs dim_pemda
conn = psycopg.connect(DB_URL)
cur = conn.cursor()

cur.execute("SELECT pemda_id, nama_pemda, tipe_pemda, provinsi_id FROM datathon.dim_pemda ORDER BY pemda_id LIMIT 20")
print("\nSample dim_pemda (20 pertama):")
print(f"  {'pemda_id':>8}  {'tipe':<12}  {'nama_pemda'}")
for pid, nama, tipe, prov_id in cur.fetchall():
    print(f"  {pid:>8}  {tipe:<12}  {nama}")

# Check apakah ID_Pemda dari sheet ada di dim_pemda
sheet_ids = set(int(k) if isinstance(k, (int, float)) else int(str(k).lstrip('0') or '0') for k in sheet_pemda.keys())
cur.execute("SELECT pemda_id FROM datathon.dim_pemda WHERE tipe_pemda = 'Provinsi'")
db_prov_ids = {r[0] for r in cur.fetchall()}

print(f"\nID dari sheet: {sorted(sheet_ids)[:10]}...")
print(f"ID Provinsi di dim_pemda: {sorted(db_prov_ids)[:10]}...")
print(f"\nMatch: {sheet_ids & db_prov_ids}")
print(f"Di sheet tapi tidak di dim_pemda: {sheet_ids - db_prov_ids}")
print(f"Di dim_pemda tapi tidak di sheet: {db_prov_ids - sheet_ids}")

conn.close()
