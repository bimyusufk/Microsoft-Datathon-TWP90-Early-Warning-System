"""
Seed dimension tables for the Datathon PostgreSQL database.

Populates:
- dim_provinsi
- dim_waktu
- dim_sektor_produktif
- dim_kualitas_pembiayaan
- dim_pemda (all distinct pemda IDs from TKD source, province-level rows mapped when possible)

This script is idempotent: it uses INSERT ... ON CONFLICT DO NOTHING/UPDATE.
"""

from __future__ import annotations

import re
from collections import OrderedDict
from datetime import date, datetime
from typing import Iterable

import psycopg
from openpyxl import load_workbook

DB_URL = "postgresql://postgres:root@localhost:5432/datathon"
XLSX_PATH = r"c:\Users\bimyu\Documents\Projects\DatathonMETC\data prep datathon dicoding.xlsx"

# ------------------------- helpers -------------------------

def normalize_province(name: str | None) -> str | None:
    if name is None:
        return None
    s = str(name).strip()
    s = re.sub(r"^(provinsi|kab\.|kabupaten|kota)\s+", "", s, flags=re.I)
    s = re.sub(r"\s+", " ", s)
    return s


def normalize_province_std(name: str | None) -> str | None:
    s = normalize_province(name)
    return s.lower() if s else None


def month_iter(start_ym: str, end_ym: str) -> Iterable[date]:
    start_y, start_m = map(int, start_ym.split("-"))
    end_y, end_m = map(int, end_ym.split("-"))
    y, m = start_y, start_m
    while (y, m) <= (end_y, end_m):
        yield date(y, m, 1)
        m += 1
        if m > 12:
            m = 1
            y += 1


def infer_pemda_type(nama_pemda: str) -> str:
    s = nama_pemda.strip().lower()
    if s.startswith("provinsi"):
        return "Provinsi"
    if s.startswith("kab.") or s.startswith("kabupaten"):
        return "Kabupaten"
    if s.startswith("kota"):
        return "Kota"
    return "Lainnya"


# ------------------------- source extraction -------------------------

wb = load_workbook(XLSX_PATH, data_only=True, read_only=True)

# Provinces: 32 from workbook + 6 Papua provinces retained as inactive master rows.
# These are the 32 provinces present in the workbook (Papua excluded from modeling).
source_provinces = [
    "Banten",
    "Dki Jakarta",
    "Jawa Barat",
    "Jawa Tengah",
    "Daerah Istimewa Yogyakarta",
    "Jawa Timur",
    "Aceh",
    "Sumatera Utara",
    "Sumatera Barat",
    "Riau",
    "Kepulauan Riau",
    "Kepulauan Bangka Belitung",
    "Jambi",
    "Sumatera Selatan",
    "Bengkulu",
    "Lampung",
    "Kalimantan Barat",
    "Kalimantan Tengah",
    "Kalimantan Utara",
    "Kalimantan Timur",
    "Kalimantan Selatan",
    "Sulawesi Utara",
    "Gorontalo",
    "Sulawesi Tengah",
    "Sulawesi Barat",
    "Sulawesi Selatan",
    "Sulawesi Tenggara",
    "Bali",
    "Nusa Tenggara Barat",
    "Nusa Tenggara Timur",
    "Maluku Utara",
    "Maluku",
]

papua_provinces = [
    "Papua",
    "Papua Barat",
    "Papua Barat Daya",
    "Papua Selatan",
    "Papua Tengah",
    "Papua Pegunungan",
]

all_provinces = source_provinces + papua_provinces

# Sector dimension from Tab 7
ws_sector = wb["Tab 7 (Penyaluran ke Sektor Pro"]
sectors = []
for row in ws_sector.iter_rows(min_row=3, values_only=True):
    if not row or row[0] in (None, ""):
        continue
    label = str(row[0]).strip()
    if label.lower() in {"total", "nasional"}:
        continue
    sectors.append(label)
sectors = list(OrderedDict.fromkeys(sectors))

# Quality dimension from Tab 10
ws_quality = wb["Tab 10 (Kualitas Pembiayaan)"]
qualities = []
for row in ws_quality.iter_rows(min_row=3, values_only=True):
    if not row or row[0] in (None, ""):
        continue
    label = str(row[0]).strip()
    if label.lower() in {"total", "nasional"}:
        continue
    qualities.append(label)
qualities = list(OrderedDict.fromkeys(qualities))

# Time dimension 2019-01 .. 2026-12
waktu_rows = []
for d in month_iter("2019-01", "2026-12"):
    tahun = d.year
    bulan = d.month
    kuartal = ((bulan - 1) // 3) + 1
    semester = 1 if bulan <= 6 else 2
    nama_bulan_id = [
        "Januari", "Februari", "Maret", "April", "Mei", "Juni",
        "Juli", "Agustus", "September", "Oktober", "November", "Desember"
    ][bulan - 1]
    waktu_rows.append((d, tahun, kuartal, bulan, nama_bulan_id, semester))

# Pemda dimension from TKD sheet
ws_tkd = wb["Realisasi TKD pertahun"]
rows = list(ws_tkd.iter_rows(values_only=True))
header = rows[0] if rows else []
col_idx = {str(v).strip(): i for i, v in enumerate(header) if v is not None}
id_idx = col_idx.get("ID_Pemda")
nama_idx = col_idx.get("Nama_Pemda")

pemda_seen = OrderedDict()
if id_idx is not None and nama_idx is not None:
    for row in rows[1:]:
        if not row or len(row) <= max(id_idx, nama_idx):
            continue
        pemda_id = row[id_idx]
        nama_pemda = row[nama_idx]
        if pemda_id in (None, "") or nama_pemda in (None, ""):
            continue
        try:
            pemda_id_int = int(float(pemda_id))
        except Exception:
            continue
        nama_pemda_str = str(nama_pemda).strip()
        tipe = infer_pemda_type(nama_pemda_str)
        pemda_seen[pemda_id_int] = (pemda_id_int, nama_pemda_str, tipe)

# build province name lookup by normalized label
prov_lookup = {normalize_province_std(p): p for p in all_provinces}

# map province-level pemda to provinsi_id later after inserting provinces
province_level_pemda = []
for pemda_id, nama_pemda, tipe in pemda_seen.values():
    if tipe == "Provinsi":
        province_level_pemda.append((pemda_id, nama_pemda, tipe))

wb.close()

# ------------------------- database seeding -------------------------

with psycopg.connect(DB_URL, autocommit=False) as conn:
    with conn.cursor() as cur:
        print("Seeding dim_provinsi...")
        # Insert 38 master provinces. Papua provinces are inactive for modeling.
        for prov in all_provinces:
            prov_std = normalize_province_std(prov)
            is_aktif = prov not in papua_provinces
            cur.execute(
                """
                INSERT INTO datathon.dim_provinsi (kode_bps, nama_provinsi, nama_provinsi_std, pulau, is_aktif)
                VALUES (%s, %s, %s, %s, %s)
                ON CONFLICT (nama_provinsi) DO UPDATE
                SET nama_provinsi_std = EXCLUDED.nama_provinsi_std,
                    pulau = EXCLUDED.pulau,
                    is_aktif = EXCLUDED.is_aktif
                """,
                (None, prov, prov_std, None, is_aktif),
            )

        print("Seeding dim_waktu...")
        for d, tahun, kuartal, bulan, nama_bulan_id, semester in waktu_rows:
            cur.execute(
                """
                INSERT INTO datathon.dim_waktu (tanggal, tahun, kuartal, bulan, nama_bulan_id, semester)
                VALUES (%s, %s, %s, %s, %s, %s)
                ON CONFLICT (tanggal) DO UPDATE
                SET tahun = EXCLUDED.tahun,
                    kuartal = EXCLUDED.kuartal,
                    bulan = EXCLUDED.bulan,
                    nama_bulan_id = EXCLUDED.nama_bulan_id,
                    semester = EXCLUDED.semester
                """,
                (d, tahun, kuartal, bulan, nama_bulan_id, semester),
            )

        print("Seeding dim_sektor_produktif...")
        for sector in sectors:
            cur.execute(
                """
                INSERT INTO datathon.dim_sektor_produktif (nama_sektor, kategori)
                VALUES (%s, %s)
                ON CONFLICT (nama_sektor) DO NOTHING
                """,
                (sector, None),
            )

        print("Seeding dim_kualitas_pembiayaan...")
        for quality in qualities:
            cur.execute(
                """
                INSERT INTO datathon.dim_kualitas_pembiayaan (nama_kualitas, parent_id, level_hirarki)
                VALUES (%s, %s, %s)
                ON CONFLICT (nama_kualitas) DO NOTHING
                """,
                (quality, None, None),
            )

        print("Seeding dim_pemda (all distinct pemda from TKD)...")
        # We store all distinct pemda ids. Province-level rows are mapped to provinsi_id;
        # kabupaten/kota rows are kept with NULL provinsi_id for now, to be mapped later.
        for pemda_id, nama_pemda, tipe in pemda_seen.values():
            provinsi_id = None
            if tipe == "Provinsi":
                prov_std = normalize_province_std(nama_pemda.replace("Provinsi", "", 1).strip())
                # match the normalized province name against inserted dim_provinsi
                cur.execute(
                    "SELECT provinsi_id FROM datathon.dim_provinsi WHERE nama_provinsi_std = %s",
                    (prov_std,),
                )
                hit = cur.fetchone()
                provinsi_id = hit[0] if hit else None

            cur.execute(
                """
                INSERT INTO datathon.dim_pemda (pemda_id, nama_pemda, tipe_pemda, provinsi_id)
                VALUES (%s, %s, %s, %s)
                ON CONFLICT (pemda_id) DO UPDATE
                SET nama_pemda = EXCLUDED.nama_pemda,
                    tipe_pemda = EXCLUDED.tipe_pemda,
                    provinsi_id = EXCLUDED.provinsi_id
                """,
                (pemda_id, nama_pemda, tipe, provinsi_id),
            )

        conn.commit()

        # Summary counts
        print("\nVerifying seed results...")
        checks = [
            ("dim_provinsi", "SELECT COUNT(*) FROM datathon.dim_provinsi"),
            ("dim_waktu", "SELECT COUNT(*) FROM datathon.dim_waktu"),
            ("dim_sektor_produktif", "SELECT COUNT(*) FROM datathon.dim_sektor_produktif"),
            ("dim_kualitas_pembiayaan", "SELECT COUNT(*) FROM datathon.dim_kualitas_pembiayaan"),
            ("dim_pemda", "SELECT COUNT(*) FROM datathon.dim_pemda"),
            ("dim_pemda_provinsi_linked", "SELECT COUNT(*) FROM datathon.dim_pemda WHERE provinsi_id IS NOT NULL"),
            ("dim_provinsi_active", "SELECT COUNT(*) FROM datathon.dim_provinsi WHERE is_aktif = TRUE"),
        ]
        for label, sql in checks:
            cur.execute(sql)
            print(f"{label} = {cur.fetchone()[0]}")

print("\nDone.")
