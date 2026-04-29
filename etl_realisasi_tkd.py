"""
ETL script: Sheet "Realisasi TKD pertahun" → datathon.fact_realisasi_tkd (PostgreSQL)

Transformasi: WIDE → LONG (unpivot 14 komponen TKD per baris sumber)

Struktur sheet (row 1 = header):
  col[0]  Tahun
  col[1]  Status_Data   ('Realisasi' | 'Alokasi')
  col[2]  ID_Pemda      string '0100' → int 100
  col[3]  Nama_Pemda
  col[4..17]  14 komponen TKD (nilai string ribuan titik atau float)

Komponen TKD (col index 4–17):
  DBH_Pajak, DBH_SDA, DBH_Lainnya, DBH_KB_LB,
  DAU_Block_Grant, DAU_Earmark, DAK_Fisik, DAK_Nonfisik,
  Hibah_Daerah, Insentif_Fiskal, Dana_Keistimewaan,
  Dana_Otsus, DTI_Papua, Dana_Desa
"""

import psycopg
from openpyxl import load_workbook

DB_URL    = "postgresql://postgres:root@localhost:5432/datathon"
XLSX_PATH = r"c:\Users\bimyu\Documents\Projects\DatathonMETC\data prep datathon dicoding.xlsx"

# Kolom komponen TKD: (col_index, nama_komponen)
KOMPONEN_COLS = [
    (4,  "DBH_Pajak"),
    (5,  "DBH_SDA"),
    (6,  "DBH_Lainnya"),
    (7,  "DBH_KB_LB"),
    (8,  "DAU_Block_Grant"),
    (9,  "DAU_Earmark"),
    (10, "DAK_Fisik"),
    (11, "DAK_Nonfisik"),
    (12, "Hibah_Daerah"),
    (13, "Insentif_Fiskal"),
    (14, "Dana_Keistimewaan"),
    (15, "Dana_Otsus"),
    (16, "DTI_Papua"),
    (17, "Dana_Desa"),
]

# ─── Helpers ──────────────────────────────────────────────────────────────────

def parse_nilai_rp(val) -> float | None:
    """
    Parse nilai TKD:
    - float/int langsung → return as-is
    - string '125.718.522.611' (titik = pemisah ribuan) → hapus titik → float
    - None / '' → None
    """
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        return float(val)
    # string
    s = str(val).strip().replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def parse_id_pemda(val) -> int | None:
    """Parse ID_Pemda: '0100' → 100, 100.0 → 100"""
    if val is None or val == "":
        return None
    try:
        return int(float(str(val)))
    except (ValueError, TypeError):
        return None


# ─── Baca XLSX ────────────────────────────────────────────────────────────────

print("Membaca sheet 'Realisasi TKD pertahun'...")
wb = load_workbook(XLSX_PATH, data_only=True, read_only=True)
ws = wb["Realisasi TKD pertahun"]
all_rows = list(ws.iter_rows(values_only=True))
wb.close()

data_rows = [
    r for r in all_rows[1:]  # skip header row
    if r and r[0] not in (None, "")
]
print(f"Baris data valid: {len(data_rows)}")
print(f"Estimasi rows setelah unpivot: {len(data_rows) * len(KOMPONEN_COLS):,}")

# ─── Koneksi ──────────────────────────────────────────────────────────────────

print("\nKoneksi ke PostgreSQL...")
conn = psycopg.connect(DB_URL)
cur  = conn.cursor()

# Verify fact_realisasi_tkd exists and is empty-ish
cur.execute("SELECT COUNT(*) FROM datathon.fact_realisasi_tkd")
existing = cur.fetchone()[0]
print(f"Rows saat ini di fact_realisasi_tkd: {existing:,}")

# ─── UNPIVOT & UPSERT ────────────────────────────────────────────────────────

upsert_sql = """
INSERT INTO datathon.fact_realisasi_tkd
    (pemda_id, tahun, status_data, komponen_tkd, nilai_rp, source_sheet)
VALUES (%s, %s, %s, %s, %s, %s)
ON CONFLICT (pemda_id, tahun, status_data, komponen_tkd) DO UPDATE SET
    nilai_rp     = EXCLUDED.nilai_rp,
    loaded_at    = CURRENT_TIMESTAMP
"""

inserted  = 0
skipped   = 0
no_pemda  = []
BATCH_SIZE = 500

print("\nMemproses unpivot dan upsert...")

batch = []

def flush_batch():
    global inserted
    if batch:
        cur.executemany(upsert_sql, batch)
        inserted += len(batch)
        batch.clear()

for i, row in enumerate(data_rows):
    # Parse identitas
    tahun = None
    try:
        tahun = int(float(str(row[0])))
    except:
        skipped += 1
        continue

    status_data = str(row[1]).strip() if row[1] not in (None, "") else None
    pemda_id    = parse_id_pemda(row[2])

    if status_data is None or pemda_id is None:
        skipped += 1
        continue

    # Unpivot: 1 baris sumber → N baris (satu per komponen)
    for col_idx, komponen in KOMPONEN_COLS:
        raw_val = row[col_idx] if len(row) > col_idx else None
        nilai   = parse_nilai_rp(raw_val)
        # Sertakan nilai 0.0, skip hanya jika None
        if nilai is None:
            continue

        batch.append((
            pemda_id,
            tahun,
            status_data,
            komponen,
            nilai,
            "Realisasi TKD pertahun",
        ))

    # Flush setiap BATCH_SIZE sumber rows
    if (i + 1) % BATCH_SIZE == 0:
        flush_batch()
        conn.commit()
        print(f"  [{i+1:>5}/{len(data_rows)}] rows sumber diproses, {inserted:>7,} facts di-insert...", end="\r")

# Flush sisa
flush_batch()
conn.commit()

print(f"\n\n✅ Selesai!")
print(f"   Rows sumber diproses : {len(data_rows) - skipped:,}")
print(f"   Rows sumber di-skip  : {skipped:,}")
print(f"   Facts upserted       : {inserted:,}")

# ─── Validasi ──────────────────────────────────────────────────────────────────

print("\n--- Validasi: per (tahun, status_data) ---")
cur.execute("""
    SELECT tahun, status_data,
           COUNT(DISTINCT pemda_id) AS n_pemda,
           COUNT(*)                 AS n_rows,
           ROUND(SUM(nilai_rp) / 1e12, 2) AS total_triliun_rp
    FROM datathon.fact_realisasi_tkd
    GROUP BY tahun, status_data
    ORDER BY tahun, status_data
""")
print(f"  {'Tahun':<6} {'Status':<12} {'N Pemda':>8} {'N Rows':>10} {'Total (T Rp)':>14}")
print("  " + "-" * 56)
for thn, status, n_pemda, n_rows, total in cur.fetchall():
    print(f"  {thn:<6} {status:<12} {n_pemda:>8,} {n_rows:>10,} {total:>14,.2f}")

cur.execute("SELECT COUNT(*) FROM datathon.fact_realisasi_tkd")
total_rows = cur.fetchone()[0]
print(f"\n  Total rows: {total_rows:,}")

# Sample join ke dim_pemda
print("\n--- Sample 5 rows (join ke dim_pemda) ---")
cur.execute("""
    SELECT f.pemda_id, p.nama_pemda, f.tahun, f.status_data,
           f.komponen_tkd, ROUND(f.nilai_rp) AS nilai_rp
    FROM datathon.fact_realisasi_tkd f
    JOIN datathon.dim_pemda p USING (pemda_id)
    WHERE f.tahun = 2024 AND f.status_data = 'Realisasi'
      AND f.komponen_tkd = 'Dana_Desa'
    ORDER BY f.pemda_id
    LIMIT 5
""")
print(f"  {'pemda_id':>9}  {'nama_pemda':<30}  {'Thn'}  {'Status':<12}  {'Komponen':<20}  {'Nilai (Rp)':>18}")
print("  " + "-" * 104)
for pid, nama, thn, status, komp, nilai in cur.fetchall():
    print(f"  {pid:>9}  {nama:<30}  {thn}  {status:<12}  {komp:<20}  {int(nilai):>18,}")

conn.close()
