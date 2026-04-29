"""
Migrasi bersih fact_perbankan:
  1. Rename kolom pemda_id → ID_Pemda (sesuai permintaan)
  2. TRUNCATE tabel
  3. Re-load 192 baris dari sheet "Perbankan" menggunakan ID_Pemda langsung
  4. Verifikasi FK relation ke dim_pemda
"""
import psycopg
from openpyxl import load_workbook

DB_URL    = "postgresql://postgres:root@localhost:5432/datathon"
XLSX_PATH = r"c:\Users\bimyu\Documents\Projects\DatathonMETC\data prep datathon dicoding.xlsx"

# Indeks kolom sheet Perbankan
COL_ID_PEMDA     = 0
COL_PROVINSI     = 1
COL_TAHUN        = 2
COL_SNAPSHOT     = 3
COL_KREDIT_TOTAL = 4
COL_DPK          = 5
COL_LDR          = 6
COL_ZONA_LDR     = 7
COL_NPL_MILIAR   = 8
COL_NPL_RATIO    = 9
COL_KREDIT_UMKM  = 10
COL_RASIO_UMKM   = 11
COL_UMKM_PER_KC  = 12
COL_JML_KC       = 13

def safe_float(val):
    if val is None or val == "": return None
    try: return float(val)
    except: return None

def safe_int(val):
    f = safe_float(val)
    return int(f) if f is not None else None

# ─── Koneksi ──────────────────────────────────────────────────────────────────
conn = psycopg.connect(DB_URL)
cur  = conn.cursor()

# ─── STEP 1: Rapikan skema ────────────────────────────────────────────────────
print("=== STEP 1: Rapikan skema fact_perbankan ===")

# Cek kolom saat ini
cur.execute("""
    SELECT column_name FROM information_schema.columns
    WHERE table_schema = 'datathon' AND table_name = 'fact_perbankan'
""")
existing = {r[0] for r in cur.fetchall()}
print(f"  Kolom saat ini: {sorted(existing)}")

# Drop PK lama
cur.execute("ALTER TABLE datathon.fact_perbankan DROP CONSTRAINT IF EXISTS fact_perbankan_pkey")
cur.execute("ALTER TABLE datathon.fact_perbankan DROP CONSTRAINT IF EXISTS fk_perbankan_pemda")

# Rename pemda_id → ID_Pemda jika perlu, atau drop dan re-add
if 'pemda_id' in existing:
    cur.execute('ALTER TABLE datathon.fact_perbankan RENAME COLUMN pemda_id TO "ID_Pemda"')
    print('  Renamed pemda_id → "ID_Pemda"')
elif '"ID_Pemda"' not in existing and 'ID_Pemda' not in existing:
    cur.execute("""
        ALTER TABLE datathon.fact_perbankan
            ADD COLUMN "ID_Pemda" INTEGER NOT NULL DEFAULT 0
    """)
    cur.execute('ALTER TABLE datathon.fact_perbankan ALTER COLUMN "ID_Pemda" DROP DEFAULT')
    print('  Added new "ID_Pemda" column')

# Add FK ke dim_pemda
cur.execute("""
    ALTER TABLE datathon.fact_perbankan
        ADD CONSTRAINT fk_perbankan_pemda
        FOREIGN KEY ("ID_Pemda") REFERENCES datathon.dim_pemda(pemda_id)
""")
# Add PK baru
cur.execute("""
    ALTER TABLE datathon.fact_perbankan
        ADD CONSTRAINT fact_perbankan_pkey
        PRIMARY KEY ("ID_Pemda", tahun, snapshot_label)
""")
conn.commit()
print('  FK → dim_pemda dan PK ("ID_Pemda", tahun, snapshot_label) berhasil dibuat')

# ─── STEP 2: Kosongkan tabel ─────────────────────────────────────────────────
print("\n=== STEP 2: Truncate fact_perbankan ===")
cur.execute("TRUNCATE TABLE datathon.fact_perbankan")
conn.commit()
cur.execute("SELECT COUNT(*) FROM datathon.fact_perbankan")
print(f"  Row count setelah truncate: {cur.fetchone()[0]}")

# Check skema akhir
print("\n=== Skema fact_perbankan (final) ===")
cur.execute("""
    SELECT column_name, data_type, is_nullable
    FROM information_schema.columns
    WHERE table_schema = 'datathon' AND table_name = 'fact_perbankan'
    ORDER BY ordinal_position
""")
for col, dtype, nullable in cur.fetchall():
    print(f"  {col:<30} {dtype:<25} {'NULL' if nullable == 'YES' else 'NOT NULL'}")

# ─── STEP 3: Load data dari XLSX ─────────────────────────────────────────────
print("\n=== STEP 3: Load data dari sheet Perbankan ===")

wb = load_workbook(XLSX_PATH, data_only=True, read_only=True)
ws = wb["Perbankan"]
all_rows = list(ws.iter_rows(values_only=True))
wb.close()

data_rows = [
    r for r in all_rows[2:]  # skip 2 header rows
    if r and r[COL_ID_PEMDA] not in (None, "")
]
print(f"  Baris valid dari XLSX: {len(data_rows)}")

upsert_sql = """
INSERT INTO datathon.fact_perbankan (
    "ID_Pemda", tahun, snapshot_label,
    kredit_total_miliar, dpk_miliar, ldr_pct,
    zona_ldr, npl_miliar, npl_ratio,
    kredit_umkm_miliar, rasio_umkm, umkm_per_kc_miliar,
    jml_kc_bank, source_sheet
) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
ON CONFLICT ("ID_Pemda", tahun, snapshot_label) DO UPDATE SET
    kredit_total_miliar = EXCLUDED.kredit_total_miliar,
    dpk_miliar          = EXCLUDED.dpk_miliar,
    ldr_pct             = EXCLUDED.ldr_pct,
    zona_ldr            = EXCLUDED.zona_ldr,
    npl_miliar          = EXCLUDED.npl_miliar,
    npl_ratio           = EXCLUDED.npl_ratio,
    kredit_umkm_miliar  = EXCLUDED.kredit_umkm_miliar,
    rasio_umkm          = EXCLUDED.rasio_umkm,
    umkm_per_kc_miliar  = EXCLUDED.umkm_per_kc_miliar,
    jml_kc_bank         = EXCLUDED.jml_kc_bank,
    loaded_at           = CURRENT_TIMESTAMP
"""

inserted = 0
skipped  = 0

for r in data_rows:
    try:
        id_pemda = int(float(str(r[COL_ID_PEMDA])))
    except:
        skipped += 1; continue

    tahun = safe_int(r[COL_TAHUN])
    snap  = str(r[COL_SNAPSHOT]).strip() if r[COL_SNAPSHOT] else None
    if tahun is None or snap is None:
        skipped += 1; continue

    cur.execute(upsert_sql, (
        id_pemda, tahun, snap,
        safe_float(r[COL_KREDIT_TOTAL]), safe_float(r[COL_DPK]), safe_float(r[COL_LDR]),
        str(r[COL_ZONA_LDR]).strip() if r[COL_ZONA_LDR] else None,
        safe_float(r[COL_NPL_MILIAR]), safe_float(r[COL_NPL_RATIO]),
        safe_float(r[COL_KREDIT_UMKM]), safe_float(r[COL_RASIO_UMKM]),
        safe_float(r[COL_UMKM_PER_KC]), safe_int(r[COL_JML_KC]),
        "Perbankan",
    ))
    inserted += 1

conn.commit()
print(f"\n  ✅ Rows inserted : {inserted}  |  skipped : {skipped}")

# ─── STEP 4: Validasi ────────────────────────────────────────────────────────
print("\n=== STEP 4: Validasi ===")

# Row count per snapshot
cur.execute("""
    SELECT tahun, snapshot_label, COUNT(*) AS n
    FROM datathon.fact_perbankan
    GROUP BY tahun, snapshot_label
    ORDER BY tahun, snapshot_label
""")
print(f"  {'Tahun':<8} {'Snapshot':<10} {'N':>4}")
print("  " + "-" * 26)
for t, s, n in cur.fetchall():
    print(f"  {t:<8} {s:<10} {n:>4}")

cur.execute("SELECT COUNT(*) FROM datathon.fact_perbankan")
print(f"\n  Total rows: {cur.fetchone()[0]}")

# Validasi JOIN ke dim_pemda
print("\n=== Join fact_perbankan ↔ dim_pemda (5 sample) ===")
cur.execute("""
    SELECT f."ID_Pemda", p.nama_pemda, f.tahun, f.snapshot_label,
           f.kredit_total_miliar, f.ldr_pct, f.zona_ldr, f.npl_ratio, f.jml_kc_bank
    FROM datathon.fact_perbankan f
    JOIN datathon.dim_pemda p ON p.pemda_id = f."ID_Pemda"
    ORDER BY f.tahun, f."ID_Pemda"
    LIMIT 5
""")
print(f"  {'ID_Pemda':>9}  {'nama_pemda':<30}  {'Thn'}  {'Snap':<8}  {'Kredit(M)':>12}  {'LDR':>8}  {'Zona':<8}  {'NPL%':>7}  {'KC':>5}")
print("  " + "-" * 105)
for row in cur.fetchall():
    pid, nama, thn, snap, kredit, ldr, zona, npl, kc = row
    print(f"  {pid:>9}  {nama:<30}  {thn}  {snap:<8}  {kredit:>12.2f}  {ldr:>8.4f}  {zona:<8}  {npl:>7.4f}  {kc:>5}")

conn.close()
print("\n✅ Selesai!")
